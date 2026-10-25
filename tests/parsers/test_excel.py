"""Excel parser against generated workbooks: every truth fact must come back at its cell range."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from fcrag.datagen import excel as gen
from fcrag.ingest.parsers.excel import ExcelParser
from fcrag.models import Element, SourceFile
from fcrag.truth import SystemTruth
from fcrag.units import display_unit


def parse(path: Path) -> list[Element]:
    sf = SourceFile(path=path, sha256="0" * 64, mime_type="x", doc_id=path.stem)
    return list(ExcelParser().parse(sf))


@pytest.mark.parametrize(("code", "version"), [("A", 1), ("A", 2), ("B", 2), ("C", 1)])
def test_every_fact_is_found_at_its_cells(
    systems: dict[str, SystemTruth], tmp_path: Path, code: str, version: int
) -> None:
    generated = gen.build_workbook(systems[code], version, tmp_path)
    elements = parse(tmp_path / generated.file)
    rows = {(e.locator["sheet"], e.locator["cells"]): e for e in elements if e.kind == "table_row"}
    assert generated.facts
    for fact in generated.facts:
        element = rows.get((fact.locator["sheet"], fact.locator["cells"]))
        assert element is not None, f"{fact.key} missing at {fact.locator}"
        assert display_unit(fact.stated_unit) in element.content, (fact.key, element.content)


def test_hidden_obsolete_row_is_skipped(systems: dict[str, SystemTruth], tmp_path: Path) -> None:
    generated = gen.build_workbook(systems["A"], 2, tmp_path)
    text = "\n".join(e.content for e in parse(tmp_path / generated.file))
    assert "obsolete" not in text
    assert "Dry-side pressure drop | Value: 15" in text


def test_layouts(systems: dict[str, SystemTruth], tmp_path: Path) -> None:
    generated = gen.build_workbook(systems["A"], 2, tmp_path)
    elements = parse(tmp_path / generated.file)
    contents = [e.content for e in elements]
    # wide layout with units row: °C converted to K in System A's sheet
    assert any("Maximum coolant outlet temperature: 368.15 K" in c for c in contents)
    # long layout keeps the component caption and the merged two-row header
    row = next(c for c in contents if "Number of cells" in c)
    assert row.startswith("Fuel cell stack (A-STK) | ID: P-")
    table = next(e for e in elements if e.kind == "table" and e.locator["sheet"] == "Stack")
    assert table.data is not None and "Operating point / Value" in table.data["columns"]
    assert table.heading_path[0] == "Stack"


def test_merged_title_and_multiple_blocks(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws["A1"] = "Two tables"
    ws.merge_cells("A1:C1")
    ws.append([])
    ws.append(["Pump"])
    ws.append(["Name", "Value", "Unit"])
    ws.append(["Flow", 150, "L/min"])
    ws.append([])
    ws.append(["Blower"])
    ws.append(["Name", "Value", "Unit"])
    ws.append(["Power", 4.5, "kW"])
    wb.save(tmp_path / "t.xlsx")
    rows = [e for e in parse(tmp_path / "t.xlsx") if e.kind == "table_row"]
    assert [r.content for r in rows] == [
        "Pump | Name: Flow | Value: 150 | Unit: L/min",
        "Blower | Name: Power | Value: 4.5 | Unit: kW",
    ]
    assert rows[1].locator == {"sheet": "Data", "cells": "A9:C9"}
    assert rows[1].heading_path == ["Data", "Two tables", "Blower"]
