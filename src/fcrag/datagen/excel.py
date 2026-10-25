"""Parameter sheets (XLSX) with deliberately messy layouts, one workbook per system and version.

Realism: merged title cells, a two-row header, several tables on one sheet, a units row with
parameters as columns, a hidden row holding an obsolete value, and units that differ from the
datasheet of the same system.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fcrag.datagen.common import (
    GeneratedFile,
    convert,
    doc_versions,
    label,
    make_fact,
    value_at_version,
)
from fcrag.truth import Component, SystemTruth, TruthParameter
from fcrag.units import display_unit

BOLD = Font(bold=True)
FILL = PatternFill("solid", fgColor="DFE7EF")


@dataclass(frozen=True)
class SheetStyle:
    language: str
    unit_variants: dict[str, str]


STYLES = {
    "A": SheetStyle("en", {"degC": "K"}),
    "B": SheetStyle("de", {"kW": "W"}),
    "C": SheetStyle("de", {"kPa": "bar"}),
}

TEXT: dict[str, dict[str, Any]] = {
    "en": {
        "title": "{system} – parameter list",
        "group_param": "Parameter",
        "group_op": "Operating point",
        "cols": ["ID", "Name", "Value", "Unit", "Condition"],
        "obsolete": "(obsolete, see revision history)",
        "limits": "Operating limits",
        "unit": "Unit",
        "value": "Value",
        "system": "System level",
    },
    "de": {
        "title": "{system} – Parameterliste",
        "group_param": "Parameter",
        "group_op": "Betriebspunkt",
        "cols": ["ID", "Bezeichnung", "Wert", "Einheit", "Bedingung"],
        "obsolete": "(veraltet, siehe Änderungshistorie)",
        "limits": "Betriebsgrenzen",
        "unit": "Einheit",
        "value": "Wert",
        "system": "Systemebene",
    },
}

LIMIT_KEYS = (
    "coolant_outlet_temperature_nominal",
    "coolant_outlet_temperature_max",
    "coolant_delta_t_max",
    "anode_cathode_dp_max",
)


def _stated(param: TruthParameter, version: int, style: SheetStyle) -> tuple[float, str]:
    value = value_at_version(param, version)
    target = style.unit_variants.get(param.unit)
    return (convert(value, param.unit, target), target) if target else (value, param.unit)


class _Writer:
    def __init__(self, system: SystemTruth, version: int, out: GeneratedFile) -> None:
        self.system, self.version, self.out = system, version, out
        self.style = STYLES[system.code]
        self.t = TEXT[self.style.language]
        self.counter = 0

    def fact(
        self, ws: Worksheet, component: Component | None, param: TruthParameter, cells: str
    ) -> None:
        value, unit = _stated(param, self.version, self.style)
        locator = {"sheet": ws.title, "cells": cells}
        self.out.facts.append(
            make_fact(self.system, component, param, value, unit, self.out, locator)
        )

    def table(
        self,
        ws: Worksheet,
        top: int,
        caption: str,
        rows: Sequence[tuple[Component | None, TruthParameter]],
    ) -> int:
        """Long-format table with a two-row header. Returns the next free row."""
        ws.cell(top, 1, caption).font = BOLD
        ws.cell(top + 1, 1, self.t["group_param"]).font = BOLD
        ws.merge_cells(start_row=top + 1, start_column=1, end_row=top + 1, end_column=2)
        ws.cell(top + 1, 3, self.t["group_op"]).font = BOLD
        ws.merge_cells(start_row=top + 1, start_column=3, end_row=top + 1, end_column=5)
        for col, name in enumerate(self.t["cols"], 1):
            cell = ws.cell(top + 2, col, name)
            cell.font, cell.fill = BOLD, FILL
        r = top + 3
        for component, param in rows:
            self.counter += 1
            value, unit = _stated(param, self.version, self.style)
            cells: list[str | float] = [
                f"P-{self.counter:03d}",
                label(param, self.style.language),
                value,
                display_unit(unit),
                param.condition or "",
            ]
            for col, v in enumerate(cells, 1):
                ws.cell(r, col, v)
            self.fact(ws, component, param, f"A{r}:E{r}")
            r += 1
            # Older values stay in the sheet, hidden, the way engineers really keep them.
            older = [rev for rev in param.superseded if rev.doc_version < self.version]
            for rev in older:
                ws.cell(r, 1, f"P-{self.counter:03d}a")
                ws.cell(r, 2, f"{label(param, self.style.language)} {self.t['obsolete']}")
                ws.cell(r, 3, rev.value)
                ws.cell(r, 4, display_unit(param.unit))
                ws.row_dimensions[r].hidden = True
                r += 1
        return r + 2

    def limits_sheet(self, wb: Workbook, stack: Component) -> None:
        """Wide layout: parameters as columns, then a units row, then a values row."""
        ws = wb.create_sheet(self.t["limits"][:31])
        ws.cell(1, 1, f"{self.system.system} – {self.t['limits']}").font = BOLD
        ws.cell(3, 1, self.t["group_param"]).font = BOLD
        ws.cell(4, 1, self.t["unit"]).font = BOLD
        ws.cell(5, 1, self.t["value"]).font = BOLD
        for col, key in enumerate(LIMIT_KEYS, 2):
            param = stack.parameters[[p.key for p in stack.parameters].index(key)]
            value, unit = _stated(param, self.version, self.style)
            ws.cell(3, col, label(param, self.style.language)).alignment = Alignment(wrap_text=True)
            ws.cell(4, col, display_unit(unit))
            ws.cell(5, col, value)
            letter = get_column_letter(col)
            self.fact(ws, stack, param, f"{letter}3:{letter}5")
            ws.column_dimensions[letter].width = 22

    def bom_sheet(self, wb: Workbook) -> None:
        """Bill of materials: the only place where the cryptic CAD part names are explained."""
        from fcrag.datagen.cad import stack_parts

        ws = wb.create_sheet("BOM")
        ws.cell(1, 1, f"{self.system.system} – stack BOM").font = BOLD
        header = ["Part ID", "Description", "Assembly", "Qty", "Size (mm)"]
        for col, name in enumerate(header, 1):
            cell = ws.cell(3, col, name)
            cell.font, cell.fill = BOLD, FILL
        for r, part in enumerate(stack_parts(self.system), 4):
            size = " x ".join(f"{v:g}" for v in part.size_mm)
            values: list[str | int] = [part.name, part.description, part.assembly or "top level",
                                       part.count, size]  # fmt: skip
            for col, value in enumerate(values, 1):
                ws.cell(r, col, value)

    def build(self, path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Stack"
        title = self.t["title"].format(system=self.system.system)
        ws.cell(1, 1, f"{title} · Rev. {self.version}").font = Font(bold=True, size=13)
        ws.merge_cells("A1:E1")
        stack = next(c for c in self.system.components if c.type == "stack")
        rows = [(stack, p) for p in stack.parameters if p.key not in LIMIT_KEYS]
        self.table(ws, 3, f"{stack.name} ({stack.id})", rows)

        bop = wb.create_sheet("BoP")
        top = 1
        for component in (c for c in self.system.components if c.type != "stack"):
            caption = f"{component.name} ({component.id})"
            top = self.table(bop, top, caption, [(component, p) for p in component.parameters])

        system_ws = wb.create_sheet(self.t["system"][:31])
        self.table(system_ws, 1, self.system.system, [(None, p) for p in self.system.parameters])
        self.limits_sheet(wb, stack)
        self.bom_sheet(wb)
        for sheet in wb.worksheets:
            for letter, width in zip("ABCDE", (10, 46, 12, 10, 40), strict=True):
                sheet.column_dimensions[letter].width = max(
                    width, sheet.column_dimensions[letter].width or 0
                )
        wb.save(path)


def build_workbook(system: SystemTruth, version: int, out_dir: Path) -> GeneratedFile:
    style = STYLES[system.code]
    doc_id = f"ps-system-{system.code.lower()}"
    t = TEXT[style.language]
    out = GeneratedFile(
        file=f"{doc_id}_v{version}.xlsx",
        doc_id=doc_id,
        doc_version=version,
        title=t["title"].format(system=system.system),
        system=system.system,
        doc_type="parameter_sheet",
        language=style.language,
    )
    _Writer(system, version, out).build(out_dir / out.file)
    return out


def build_all(system: SystemTruth, out_dir: Path) -> list[GeneratedFile]:
    return [build_workbook(system, v, out_dir) for v in range(1, doc_versions(system) + 1)]
