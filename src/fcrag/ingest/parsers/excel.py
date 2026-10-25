"""Excel parser: table blocks, merged cells, multi-row headers, units rows, hidden rows.

Each sheet is cut into blocks of consecutive non-empty rows. A block yields one ``table``
element (Markdown, for overview questions) and one ``table_row`` fact per data row, each cited
by sheet and cell range. Two layouts are recognised:

- long: a header (one or more rows) followed by one parameter per row;
- wide: a row starting with "Unit"/"Einheit" under a header row, one parameter per column.
"""

from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fcrag.ingest.parsers import register
from fcrag.models import Element, SourceFile

UNIT_LABELS = {"unit", "units", "einheit", "einheiten"}
VALUE_LABELS = {"value", "values", "wert", "werte"}


@dataclass
class Row:
    number: int
    """1-based worksheet row, for citations."""
    cells: list[Any]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    return " ".join(str(value).split())


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def read_grid(ws: Worksheet) -> list[Row]:
    """Visible rows with merged ranges filled from their top-left cell."""
    fill: dict[tuple[int, int], Any] = {}
    for merged in ws.merged_cells.ranges:
        top_left = ws.cell(merged.min_row, merged.min_col).value
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                fill[(r, c)] = top_left
    hidden_cols = {
        i for i in range(1, ws.max_column + 1)
        if ws.column_dimensions[get_column_letter(i)].hidden
    }  # fmt: skip
    rows = []
    for r in range(1, ws.max_row + 1):
        if ws.row_dimensions[r].hidden:
            continue
        cells = [
            fill.get((r, c), ws.cell(r, c).value)
            for c in range(1, ws.max_column + 1)
            if c not in hidden_cols
        ]
        rows.append(Row(r, cells))
    return rows


def split_blocks(rows: list[Row]) -> list[list[Row]]:
    blocks: list[list[Row]] = []
    current: list[Row] = []
    for row in rows:
        if any(_text(v) for v in row.cells):
            current.append(row)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def _is_title(row: Row) -> bool:
    """A caption row: one distinct non-empty value (a merged title fills several cells)."""
    values = {_text(v) for v in row.cells if _text(v)}
    return len(values) == 1 and not any(_is_number(v) for v in row.cells)


def _width(block: list[Row]) -> int:
    return max(max((i + 1 for i, v in enumerate(r.cells) if _text(v)), default=0) for r in block)


def _cells(first: Row, last: Row, width: int, col: int | None = None) -> str:
    if col is not None:
        letter = get_column_letter(col + 1)
        return f"{letter}{first.number}:{letter}{last.number}"
    return f"A{first.number}:{get_column_letter(width)}{last.number}"


def _combine_headers(header_rows: list[Row], width: int) -> list[str]:
    names = []
    for c in range(width):
        parts: list[str] = []
        for row in header_rows:
            text = _text(row.cells[c]) if c < len(row.cells) else ""
            if text and text not in parts:
                parts.append(text)
        names.append(" / ".join(parts) or f"Column {c + 1}")
    return names


def _markdown(header: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(lines)


class ExcelParser:
    supported_types = frozenset(
        {".xlsx", ".xlsm", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    )

    def parse(self, file: SourceFile) -> Iterator[Element]:
        wb = load_workbook(file.path, data_only=True)  # cached formula results, not formulas
        try:
            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    continue
                yield from self._sheet(ws)
        finally:
            wb.close()

    def _sheet(self, ws: Worksheet) -> Iterator[Element]:
        sheet_title: str | None = None
        for block in split_blocks(read_grid(ws)):
            captions = []
            while block and _is_title(block[0]):
                captions.append(next(_text(v) for v in block[0].cells if _text(v)))
                block = block[1:]
            if not block:  # a block of titles only, e.g. the sheet heading
                sheet_title = sheet_title or (captions[0] if captions else None)
                continue
            path = [p for p in (ws.title, sheet_title, *captions) if p]
            caption = captions[-1] if captions else (sheet_title or ws.title)
            units_at = next(
                (i for i, r in enumerate(block) if _text(r.cells[0]).lower() in UNIT_LABELS), None
            )
            if units_at is not None and units_at > 0:
                yield from self._wide(ws, block, units_at, caption, path)
            else:
                yield from self._long(ws, block, caption, path)

    def _long(
        self, ws: Worksheet, block: list[Row], caption: str, path: list[str]
    ) -> Iterator[Element]:
        width = _width(block)
        first_data = next((i for i, r in enumerate(block) if any(map(_is_number, r.cells))), 0)
        header_rows, data_rows = block[:first_data], block[first_data:]
        header = _combine_headers(header_rows, width) if header_rows else [
            f"Column {c + 1}" for c in range(width)
        ]  # fmt: skip
        leaf = [h.split(" / ")[-1] for h in header]
        table: list[list[str]] = []
        for row in data_rows:
            values = [_text(row.cells[c]) if c < len(row.cells) else "" for c in range(width)]
            table.append(values)
            pairs = [f"{name}: {v}" for name, v in zip(leaf, values, strict=True) if v]
            if len(pairs) < 2:
                continue
            yield Element(
                kind="table_row",
                content=f"{caption} | " + " | ".join(pairs),
                heading_path=path,
                locator={"sheet": ws.title, "cells": _cells(row, row, width)},
                data={"caption": caption, "columns": leaf, "values": values},
            )
        if table:
            yield Element(
                kind="table",
                content=f"{caption}\n\n{_markdown(header, table)}",
                heading_path=path,
                locator={"sheet": ws.title, "cells": _cells(block[0], block[-1], width)},
                data={"columns": header, "rows": table},
            )

    def _wide(
        self, ws: Worksheet, block: list[Row], units_at: int, caption: str, path: list[str]
    ) -> Iterator[Element]:
        header_row, units_row = block[units_at - 1], block[units_at]
        value_rows = block[units_at + 1 :]
        width = _width(block)
        columns: list[list[str]] = []
        for c in range(1, width):
            name = _text(header_row.cells[c])
            unit = _text(units_row.cells[c])
            for row in value_rows:
                value = _text(row.cells[c])
                if not (name and value):
                    continue
                row_label = _text(row.cells[0])
                prefix = "" if row_label.lower() in VALUE_LABELS else f"{row_label} "
                columns.append([name, value, unit])
                yield Element(
                    kind="table_row",
                    content=f"{caption} | {prefix}{name}: {value} {unit}".rstrip(),
                    heading_path=path,
                    locator={"sheet": ws.title, "cells": _cells(header_row, row, width, col=c)},
                    data={"caption": caption, "parameter": name, "value": value, "unit": unit},
                )
        if columns:
            yield Element(
                kind="table",
                content=f"{caption}\n\n{_markdown(['Parameter', 'Value', 'Unit'], columns)}",
                heading_path=path,
                locator={"sheet": ws.title, "cells": _cells(block[0], block[-1], width)},
                data={"columns": ["Parameter", "Value", "Unit"], "rows": columns},
            )


register(ExcelParser())
